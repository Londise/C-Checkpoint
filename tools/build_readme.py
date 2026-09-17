from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\londa\Downloads\CP4.Catalogo-main\README.docx"


def set_run_font(run, name="Aptos", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True, color=(0, 0, 0))
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas", size=9.5)
    return paragraph


def add_endpoint_table(document):
    rows = [
        ("GET", "/api/produtos", "Lista produtos", "200"),
        ("GET", "/api/produtos/{id}", "Consulta produto", "200 ou 404"),
        ("POST", "/api/produtos", "Cadastra produto", "201 ou 400"),
        ("PUT", "/api/produtos/{id}", "Atualiza produto", "204, 400 ou 404"),
        ("DELETE", "/api/produtos/{id}", "Exclui produto", "204 ou 404"),
        ("GET", "/api/categorias", "Lista categorias", "200"),
        ("POST", "/api/categorias", "Cadastra categoria", "201 ou 400"),
    ]
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [Inches(0.78), Inches(1.7), Inches(2.55), Inches(1.35)]
    headers = ["Método", "Endpoint", "Objetivo", "HTTP"]

    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        set_run_font(run, size=9.5, bold=True, color=(255, 255, 255))

    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F3F6FA")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, size=9.25, bold=index == 0)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("CP4 Catalogo Guia de Execucao")
    set_run_font(title_run, size=22, bold=True, color=(0, 0, 0))

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("FIAP C Sharp Software Development  Segundo semestre de 2026")
    set_run_font(subtitle_run, size=10.5, color=(70, 70, 70))

    add_body(document, "Este README descreve como executar a solução CP4 Catalogo e como alternar a persistência entre SQL Server e Oracle. A aplicação Web consome a API por HTTP e não acessa o banco de dados ou o DbContext diretamente.")

    add_heading(document, "Identificacao do grupo")
    add_bullet(document, "Integrante identificado nos scripts: Fernando  RM 558095")
    add_bullet(document, "Turma: FIAP C Sharp Software Development  Segundo semestre de 2026")
    add_body(document, "Se houver outros integrantes, inclua os respectivos nomes e RMs nesta seção antes da entrega final.")

    add_heading(document, "Estrutura da solucao")
    add_bullet(document, "CP4.Catalogo.Data: entidades Categoria e Produto, CatalogoDbContext e configuração dos providers.")
    add_bullet(document, "CP4.Catalogo.Api: endpoints REST, DTOs, validação, Swagger e leitura do header X Database Provider.")
    add_bullet(document, "CP4.Catalogo.Web: interface MVC, HttpClient, async await e seletor de provider.")

    add_heading(document, "Pre requisitos")
    add_bullet(document, "SDK .NET 8 instalado.")
    add_bullet(document, "SQL Server ou LocalDB para executar o provider SqlServer.")
    add_bullet(document, "Oracle XE com o PDB XEPDB1 ou o ambiente Oracle institucional para executar o provider Oracle.")
    add_bullet(document, "Certificado HTTPS de desenvolvimento confiável para usar os perfis HTTPS locais.")

    add_heading(document, "Preparacao dos bancos")
    add_body(document, "SQL Server: execute scripts CP4_Fernando_RM_558095_SQLServer.sql. O script cria o database CP4_Fernando_RM_558095, as tabelas CATEGORIAS e PRODUTOS, o relacionamento e dados de exemplo.")
    add_body(document, "Oracle XE local: em uma conexão administrativa no serviço XEPDB1, execute scripts CP4_Fernando_RM_558095_Oracle_XE_Local_Admin.sql. Troque a senha de exemplo no script e use a mesma senha na configuração da API.")
    add_body(document, "Oracle institucional: conecte com o usuário autorizado e execute scripts CP4_Fernando_RM_558095_Oracle_Institucional.sql. Esse script não cria usuários.")

    add_heading(document, "Connection strings")
    add_body(document, "As connection strings ficam em CP4.Catalogo.Api appsettings.json. A configuração padrão seleciona SqlServer. Ajuste a string Oracle antes de testar esse provider. Não registre senhas reais neste documento, no repositório ou nas evidências públicas.")
    add_code(document, '"DefaultProvider": "SqlServer"')
    add_code(document, '"SqlServer": "Server=(localdb)\\MSSQLLocalDB;Database=CP4_Fernando_RM_558095;..."')
    add_code(document, '"Oracle": "User Id=CP4_FERNANDO_RM_558095;Password=SUA_SENHA;Data Source=localhost:1521/XEPDB1"')

    add_heading(document, "Como iniciar a aplicacao")
    add_body(document, "Abra CP4.Catalogo.sln no Visual Studio e configure CP4.Catalogo.Api e CP4.Catalogo.Web como projetos de inicialização múltipla. Como alternativa, use dois terminais na raiz da solução.")
    add_code(document, "dotnet run --project CP4.Catalogo.Api --launch-profile https")
    add_code(document, "dotnet run --project CP4.Catalogo.Web --launch-profile https")
    add_body(document, "Com os perfis atuais, o Swagger fica em https://localhost:7030/swagger e a interface MVC fica em https://localhost:7113. A Web usa https://localhost:7030 como endereço base da API.")

    add_heading(document, "Selecao do provider")
    add_body(document, "Na interface MVC, escolha SQL Server ou Oracle no seletor da tela Catálogo de Produtos. A Web encaminha a escolha no header X Database Provider. A API valida o valor, responde com o provider utilizado no mesmo header e cria o CatalogoDbContext com UseSqlServer ou UseOracle.")
    add_body(document, "No Swagger, informe o mesmo header em cada operação. Valores aceitos: SqlServer e Oracle. Se o header estiver ausente, a API usa o DefaultProvider de appsettings.json.")

    add_heading(document, "Endpoints obrigatorios")
    add_endpoint_table(document)

    add_heading(document, "Checklist de entrega")
    add_bullet(document, "Compile a solução e execute GET api produtos com SqlServer e Oracle.")
    add_bullet(document, "Registre as evidências do Swagger, dos dois bancos e da interface MVC mostrando o provider selecionado.")
    add_bullet(document, "Entregue CP4.Catalogo.sln, os três projetos, a pasta scripts e este README.")
    add_bullet(document, "Não inclua as pastas bin e obj no arquivo ZIP.")

    document.core_properties.title = "CP4 Catalogo Guia de Execucao"
    document.core_properties.subject = "Instruções de execução da solução CP4 Catalogo"
    document.core_properties.author = "Fernando RM 558095"
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
